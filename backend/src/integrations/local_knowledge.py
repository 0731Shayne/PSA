"""Lightweight local RAG primitives for portfolio and offline demonstrations.

This deliberately favors clarity over large-scale performance: Ollama creates
real neural embeddings, SQLite stores JSON vectors, and Python calculates
cosine similarity. The surrounding PSA permissions and citations are shared
with the production RAGFlow backend.
"""

from __future__ import annotations

import io
import json
import math
import re
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

import httpx
from docx import Document as WordDocument
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from ..config import (
    LOCAL_CHUNK_OVERLAP,
    LOCAL_CHUNK_SIZE,
    LOCAL_EMBEDDING_BASE_URL,
    LOCAL_EMBEDDING_BATCH_SIZE,
    LOCAL_EMBEDDING_MODEL,
    LOCAL_EMBEDDING_TIMEOUT_SECONDS,
    LOCAL_EMBEDDING_USE_ENV_PROXY,
    RAGFLOW_RETRIEVAL_PAGE_SIZE,
    RAGFLOW_SIMILARITY_THRESHOLD,
)
from ..db.models import KnowledgeChunk, KnowledgeDocument


LOCAL_EXTENSIONS = {".pdf", ".docx", ".pptx", ".xlsx", ".txt", ".md", ".html", ".htm"}
MAX_EXTRACTED_CHARACTERS = 2_000_000


class LocalKnowledgeError(RuntimeError):
    """A user-safe local parsing or embedding failure."""


class _TextHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        value = data.strip()
        if value:
            self.parts.append(value)


class OllamaEmbeddingClient:
    """Minimal async client for Ollama's batched `/api/embed` endpoint."""

    def __init__(self, *, transport: httpx.AsyncBaseTransport | None = None) -> None:
        self._transport = transport

    async def health(self) -> dict[str, Any]:
        payload = await self._request("GET", "/api/tags")
        models = payload.get("models", []) if isinstance(payload, dict) else []
        names = {
            str(item.get("name") or item.get("model") or "")
            for item in models
            if isinstance(item, dict)
        }
        available = any(
            name == LOCAL_EMBEDDING_MODEL
            or name.startswith(f"{LOCAL_EMBEDDING_MODEL}:")
            for name in names
        )
        return {"reachable": True, "model_available": available, "models": sorted(names)}

    async def embed(self, texts: list[str]) -> list[list[float]]:
        if not texts:
            return []
        embeddings: list[list[float]] = []
        for start in range(0, len(texts), LOCAL_EMBEDDING_BATCH_SIZE):
            batch = texts[start : start + LOCAL_EMBEDDING_BATCH_SIZE]
            payload = await self._request(
                "POST",
                "/api/embed",
                json={"model": LOCAL_EMBEDDING_MODEL, "input": batch, "truncate": True},
            )
            raw = payload.get("embeddings") if isinstance(payload, dict) else None
            if not isinstance(raw, list) or len(raw) != len(batch):
                raise LocalKnowledgeError("本地 Embedding 服务返回了无效向量")
            for vector in raw:
                if not isinstance(vector, list) or not vector:
                    raise LocalKnowledgeError("本地 Embedding 服务返回了空向量")
                try:
                    embeddings.append([float(value) for value in vector])
                except (TypeError, ValueError) as exc:
                    raise LocalKnowledgeError("本地 Embedding 向量格式无效") from exc
        dimensions = {len(vector) for vector in embeddings}
        if len(dimensions) != 1:
            raise LocalKnowledgeError("本地 Embedding 向量维度不一致")
        return embeddings

    async def _request(
        self,
        method: str,
        path: str,
        *,
        json: dict[str, Any] | None = None,
    ) -> Any:
        try:
            async with httpx.AsyncClient(
                base_url=LOCAL_EMBEDDING_BASE_URL,
                timeout=LOCAL_EMBEDDING_TIMEOUT_SECONDS,
                trust_env=LOCAL_EMBEDDING_USE_ENV_PROXY,
                transport=self._transport,
            ) as client:
                response = await client.request(method, path, json=json)
                response.raise_for_status()
                return response.json()
        except (httpx.HTTPError, ValueError) as exc:
            raise LocalKnowledgeError(
                "无法连接本地 Embedding 服务，请确认 Ollama 已启动并执行 ollama pull bge-m3"
            ) from exc


def extract_document_text(filename: str, content: bytes) -> str:
    """Extract bounded plain text from common modern course-document formats."""

    suffix = Path(filename).suffix.lower()
    if suffix not in LOCAL_EXTENSIONS:
        raise LocalKnowledgeError("本地模式支持 PDF、DOCX、PPTX、XLSX、TXT、Markdown 和 HTML")
    try:
        if suffix == ".pdf":
            text = "\n\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages)
        elif suffix == ".docx":
            document = WordDocument(io.BytesIO(content))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        elif suffix == ".pptx":
            presentation = Presentation(io.BytesIO(content))
            text = "\n\n".join(
                "\n".join(
                    str(shape.text)
                    for shape in slide.shapes
                    if hasattr(shape, "text") and str(shape.text).strip()
                )
                for slide in presentation.slides
            )
        elif suffix == ".xlsx":
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            lines: list[str] = []
            for sheet in workbook.worksheets:
                lines.append(f"# {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value) for value in row if value is not None]
                    if values:
                        lines.append("\t".join(values))
            workbook.close()
            text = "\n".join(lines)
        else:
            decoded = _decode_text(content)
            if suffix in {".html", ".htm"}:
                parser = _TextHTMLParser()
                parser.feed(decoded)
                text = "\n".join(parser.parts)
            else:
                text = decoded
    except LocalKnowledgeError:
        raise
    except Exception as exc:
        raise LocalKnowledgeError("课程文件解析失败，请确认文件没有损坏或加密") from exc
    normalized = _normalize_text(text)
    if not normalized:
        raise LocalKnowledgeError("没有从文件中提取到可检索文字")
    if len(normalized) > MAX_EXTRACTED_CHARACTERS:
        raise LocalKnowledgeError("本地演示模式单个文件最多提取 200 万字符")
    return normalized


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise LocalKnowledgeError("文本文件编码无法识别，请转换为 UTF-8")


def _normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def chunk_text(
    text: str,
    *,
    chunk_size: int = LOCAL_CHUNK_SIZE,
    overlap: int = LOCAL_CHUNK_OVERLAP,
) -> list[str]:
    """Create overlapping character chunks, preferring paragraph/sentence boundaries."""

    normalized = _normalize_text(text)
    if len(normalized) <= chunk_size:
        return [normalized] if normalized else []
    chunks: list[str] = []
    cursor = 0
    minimum_boundary = int(chunk_size * 0.6)
    while cursor < len(normalized):
        hard_end = min(cursor + chunk_size, len(normalized))
        end = hard_end
        if hard_end < len(normalized):
            window = normalized[cursor + minimum_boundary : hard_end]
            boundary = max(window.rfind("\n\n"), window.rfind("。"), window.rfind("；"))
            if boundary >= 0:
                end = cursor + minimum_boundary + boundary + 1
        chunk = normalized[cursor:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(normalized):
            break
        cursor = max(cursor + 1, end - overlap)
    return chunks


async def index_local_document(
    db: AsyncSession,
    document: KnowledgeDocument,
    *,
    client: OllamaEmbeddingClient | None = None,
) -> int:
    """Embed extracted text and atomically replace a document's local chunks."""

    text = (document.extracted_text or "").strip()
    if not text:
        raise LocalKnowledgeError("本地文档缺少可重新索引的文本")
    contents = chunk_text(text)
    vectors = await (client or OllamaEmbeddingClient()).embed(contents)
    await db.execute(delete(KnowledgeChunk).where(KnowledgeChunk.document_id == document.id))
    db.add_all(
        [
            KnowledgeChunk(
                document_id=document.id,
                position=position,
                content=content,
                embedding_json=json.dumps(vector, separators=(",", ":")),
            )
            for position, (content, vector) in enumerate(zip(contents, vectors, strict=True))
        ]
    )
    document.status = "ready"
    document.last_error = None
    await db.commit()
    await db.refresh(document)
    return len(contents)


async def retrieve_local_materials(
    db: AsyncSession,
    question: str,
    documents: list[KnowledgeDocument],
    *,
    client: OllamaEmbeddingClient | None = None,
) -> list[dict[str, Any]]:
    """Return top cosine-similar local chunks from already authorized documents."""

    if not documents:
        return []
    query_vectors = await (client or OllamaEmbeddingClient()).embed([question])
    query_vector = query_vectors[0]
    document_by_id = {document.id: document for document in documents}
    chunks = list(
        (
            await db.execute(
                select(KnowledgeChunk).where(
                    KnowledgeChunk.document_id.in_(list(document_by_id))
                )
            )
        ).scalars().all()
    )
    scored: list[tuple[float, KnowledgeChunk]] = []
    for chunk in chunks:
        try:
            vector = [float(value) for value in json.loads(chunk.embedding_json)]
        except (TypeError, ValueError, json.JSONDecodeError):
            continue
        similarity = cosine_similarity(query_vector, vector)
        if similarity >= RAGFLOW_SIMILARITY_THRESHOLD:
            scored.append((similarity, chunk))
    scored.sort(key=lambda item: item[0], reverse=True)
    sources: list[dict[str, Any]] = []
    for similarity, chunk in scored[:RAGFLOW_RETRIEVAL_PAGE_SIZE]:
        document = document_by_id[chunk.document_id]
        sources.append(
            {
                "kind": "material",
                "id": f"local-chunk-{chunk.id}",
                "document_id": document.provider_document_id,
                "document_name": document.filename,
                "content": chunk.content,
                "similarity": similarity,
                "positions": [chunk.position],
                "backend": "local",
            }
        )
    return sources


def cosine_similarity(left: list[float], right: list[float]) -> float:
    if not left or len(left) != len(right):
        return 0.0
    left_norm = math.sqrt(sum(value * value for value in left))
    right_norm = math.sqrt(sum(value * value for value in right))
    if left_norm == 0 or right_norm == 0:
        return 0.0
    return sum(a * b for a, b in zip(left, right, strict=True)) / (left_norm * right_norm)
